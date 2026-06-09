"""
Comprehensive pytest test suite for Voltík – EV wallbox planning app.

Covers: capacity_calculator, recommender, scheduler, train_models, document_generator.
Some tests are expected to FAIL due to the known root problem: the ML classifier
is trained only on zone-level features and ignores building parameters (breaker,
flats, parking, EVs), so changing those sliders does not change its output.
"""

import math
import sys
from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest

# ---------------------------------------------------------------------------
# Make sure the project's src/ is importable regardless of working directory
# ---------------------------------------------------------------------------
ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT / "src"))

# ---------------------------------------------------------------------------
# Shared fixtures / helpers
# ---------------------------------------------------------------------------

def _make_zone_row(
    grid_zone_id: str = "TS_TEST",
    reserve_margin: float = 25.0,
    grid_sensitivity: float = 0.5,
) -> pd.Series:
    """Minimal zone row for recommender tests."""
    return pd.Series({
        "grid_zone_id": grid_zone_id,
        "reserve_margin_pct_2025_synthetic": reserve_margin,
        "grid_sensitivity_index_derived": grid_sensitivity,
        "residential_index_derived": 0.5,
        "destination_activity_index_derived": 0.5,
        "no_private_parking_index_derived": 0.5,
        "reserve_capacity_kw_2025_synthetic": 100.0,
        "target_daily_charging_kwh_2030_synthetic": 50.0,
        "target_peak_charging_kw_2030_synthetic": 20.0,
        "charging_points_2026_real": 2,
        "charging_station_power_kw_2026_real": 22.0,
        "parking_segments_zps_real": 10,
        "population_model_input_derived": 500.0,
    })


def _hourly_agg_for_zone(zone_id: str = "TS_FAKE") -> pd.DataFrame:
    """Return a synthetic 24-row hourly aggregation for a single zone."""
    return pd.DataFrame({
        "grid_zone_id": [zone_id] * 24,
        "hour": list(range(24)),
        "avg_available_kw": [30.0] * 24,
        "overload_rate": [0.0] * 24,
    })


# ===========================================================================
# 1. capacity_calculator
# ===========================================================================

class TestBreakerToKw:
    def test_63A_approx_41_5_kw(self):
        from capacity_calculator import breaker_to_kw
        result = breaker_to_kw(63)
        # sqrt(3) * 400 * 63 * 0.95 / 1000 ≈ 41.5 kW
        assert abs(result - 41.5) < 0.2, f"Expected ~41.5 kW, got {result}"

    def test_returns_float(self):
        from capacity_calculator import breaker_to_kw
        assert isinstance(breaker_to_kw(63), float)


class TestCalculateWallboxCapacity:
    def _call(self, flats=24, parking=12, breaker=63, evs=5, has_pv=False, pv_kwp=0):
        from capacity_calculator import calculate_wallbox_capacity
        return calculate_wallbox_capacity(flats, parking, breaker, evs, has_pv, pv_kwp)

    def test_available_for_ev_kw_positive_standard_inputs(self):
        result = self._call(flats=24, parking=12, breaker=63, evs=5)
        assert result["available_for_ev_kw"] > 0, (
            f"Expected positive available_for_ev_kw, got {result['available_for_ev_kw']}"
        )

    def test_different_breakers_produce_different_available_kw(self):
        low = self._call(breaker=25)["available_for_ev_kw"]
        high = self._call(breaker=160)["available_for_ev_kw"]
        assert high > low, (
            f"160A breaker should give more available kW than 25A, got {high} vs {low}"
        )

    def test_warning_red_when_available_below_7kw(self):
        # 10 flats × 1.4 kW evening peak = 14 kW; breaker 25A ≈ 16.5 kW → available ≈ 2.5 kW
        result = self._call(flats=10, parking=5, breaker=25, evs=1)
        assert result["warning_level"] == "red", (
            f"Expected red warning for low available kW ({result['available_for_ev_kw']} kW), "
            f"got {result['warning_level']}"
        )

    def test_warning_green_comfortable(self):
        # 5 flats × 1.4 = 7 kW evening peak; breaker 160A ≈ 105 kW → large headroom
        result = self._call(flats=5, parking=10, breaker=160, evs=2)
        assert result["warning_level"] == "green", (
            f"Expected green warning with plenty of headroom, got {result['warning_level']}"
        )

    def test_warning_orange_load_balancing_needed(self):
        # Tune to hit orange: enough capacity but not comfortable enough for green
        # 20 flats × 1.4 = 28 kW; breaker 63A ≈ 41.5 kW → available ≈ 13.5 kW
        # max_without_balancing = floor(13.5/11) = 1 → orange
        result = self._call(flats=20, parking=10, breaker=63, evs=5)
        assert result["warning_level"] in ("orange", "green"), (
            f"Expected orange or green, got {result['warning_level']}"
        )
        # More specific: with exactly 1 wallbox without balancing it should be orange
        from capacity_calculator import calculate_wallbox_capacity
        r2 = calculate_wallbox_capacity(20, 10, 63, 5, False, 0)
        if r2["max_wallboxes_without_balancing"] < 2:
            assert r2["warning_level"] == "orange"

    def test_available_kw_is_zero_or_positive_never_negative(self):
        # Even absurdly small breaker should not produce negative available kW
        result = self._call(flats=100, parking=5, breaker=16, evs=1)
        assert result["available_for_ev_kw"] >= 0

    def test_pv_contributes_pv_peak_kw(self):
        result = self._call(has_pv=True, pv_kwp=10.0)
        assert result["pv_peak_kw"] > 0

    def test_no_pv_zero_pv_peak_kw(self):
        result = self._call(has_pv=False)
        assert result["pv_peak_kw"] == 0.0

    def test_result_keys_present(self):
        result = self._call()
        for key in [
            "total_building_capacity_kw",
            "estimated_evening_base_load_kw",
            "available_for_ev_kw",
            "max_wallboxes_without_balancing",
            "max_simultaneous_with_balancing",
            "recommended_installed_wallboxes",
            "pv_peak_kw",
            "warning_level",
        ]:
            assert key in result, f"Missing key: {key}"


# ===========================================================================
# 2. recommender
# ===========================================================================

class TestRecommender:
    def _call(
        self,
        expected_evs=5,
        available_for_ev_kw=20.0,
        reserve_margin=25.0,
        grid_sensitivity=0.5,
        ml_solution_type="residential_ac_small",
        ml_lb_type="static",
        wallboxes=4,
    ):
        from recommender import recommend_configuration
        user_input = {"expected_evs": expected_evs, "some_zone": "TS_TEST"}
        zone_row = _make_zone_row(reserve_margin=reserve_margin, grid_sensitivity=grid_sensitivity)
        capacity_result = {
            "available_for_ev_kw": available_for_ev_kw,
            "recommended_installed_wallboxes": wallboxes,
        }
        ml_prediction = {
            "solution_type": ml_solution_type,
            "load_balancing_type": ml_lb_type,
        }
        return recommend_configuration(user_input, zone_row, capacity_result, ml_prediction)

    def test_low_available_kw_forces_none_monitor(self):
        result = self._call(available_for_ev_kw=5.0)
        assert result["recommended_solution_type"] == "none_monitor", (
            f"Expected none_monitor when available < 7 kW, got {result['recommended_solution_type']}"
        )

    def test_low_available_kw_sets_load_balancing_none(self):
        result = self._call(available_for_ev_kw=5.0)
        assert result["load_balancing_type"] == "none"

    def test_high_grid_sensitivity_triggers_dynamic_load_balancing(self):
        result = self._call(grid_sensitivity=0.9, available_for_ev_kw=30.0)
        assert result["load_balancing_type"] == "dynamic", (
            f"Expected dynamic load balancing for high grid sensitivity, "
            f"got {result['load_balancing_type']}"
        )

    def test_low_grid_sensitivity_uses_static_load_balancing(self):
        result = self._call(grid_sensitivity=0.3, available_for_ev_kw=30.0)
        assert result["load_balancing_type"] == "static"

    def test_result_keys_present(self):
        result = self._call()
        for key in [
            "recommended_solution_type",
            "solution_label",
            "number_of_wallboxes",
            "recommended_total_kw",
            "load_balancing_type",
            "load_balancing_label",
            "explanation",
            "risk_score",
        ]:
            assert key in result, f"Missing key: {key}"

    def test_risk_score_bounded_0_to_100(self):
        result = self._call()
        assert 0 <= result["risk_score"] <= 100

    def test_low_reserve_margin_adds_warning_text(self):
        result = self._call(reserve_margin=5.0, available_for_ev_kw=20.0)
        assert result["risk_score"] >= 30, (
            "Low reserve margin should add risk score >= 30"
        )

    def test_none_monitor_sets_wallboxes_to_zero(self):
        result = self._call(available_for_ev_kw=3.0)
        assert result["number_of_wallboxes"] == 0

    # -----------------------------------------------------------------------
    # CRITICAL: This test is expected to FAIL.
    # The ML classifier is trained on zone-only features, so changing the breaker
    # (a building parameter) does not affect predict_solution() output.
    # The recommender therefore returns the same solution_type for both a 25A
    # and a 160A breaker when zone conditions and safety overrides stay neutral.
    # -----------------------------------------------------------------------
    def test_different_breakers_change_recommendation_KNOWN_FAIL(self):
        """
        KNOWN FAILING TEST.

        When changing breaker from 25A (weak) to 160A (strong), the resulting
        capacity is very different, so the ML prediction (solution_type) itself
        should differ — the ML model should encode the building capacity in
        its prediction, not just the zone.

        Currently fails because predict_solution() does now accept capacity_result
        but the model was trained without these features, so it ignores them and
        returns the same zone-level prediction regardless.
        """
        from capacity_calculator import calculate_wallbox_capacity
        from train_models import load_or_train_classifier, predict_solution

        model_tuple = load_or_train_classifier()
        zone_row = _make_zone_row(reserve_margin=25.0, grid_sensitivity=0.4)

        user_input = {"expected_evs": 5, "parking_spaces": 12, "flats": 24}

        cap_weak = calculate_wallbox_capacity(
            flats=24, parking_spaces=12, breaker_amps=25,
            expected_evs=5, has_pv=False, pv_kwp=0,
        )
        cap_strong = calculate_wallbox_capacity(
            flats=24, parking_spaces=12, breaker_amps=160,
            expected_evs=5, has_pv=False, pv_kwp=0,
        )

        ml_weak = predict_solution(model_tuple, zone_row, cap_weak, user_input)
        ml_strong = predict_solution(model_tuple, zone_row, cap_strong, user_input)

        # The ML prediction itself must differ: a 25A breaker leaves ~0 kW for EVs,
        # so the model should predict none_monitor; a 160A breaker leaves ~71 kW,
        # so the model should predict a real wallbox solution.
        assert ml_weak["solution_type"] != ml_strong["solution_type"], (
            "KNOWN FAIL: ML predict_solution() returns identical solution_type for "
            f"25A breaker ({ml_weak['solution_type']}) and 160A breaker "
            f"({ml_strong['solution_type']}), even though available_for_ev_kw differs "
            f"dramatically ({cap_weak['available_for_ev_kw']} kW vs "
            f"{cap_strong['available_for_ev_kw']} kW). "
            "This confirms the ML classifier was trained without building parameters."
        )


# ===========================================================================
# 3. scheduler
# ===========================================================================

class TestScheduler:
    def _plan(self, has_pv=False, pv_kwp=0.0, expected_daily_kwh=50.0, zone_id="TS_FAKE"):
        from scheduler import create_hourly_charging_plan
        hourly_agg = _hourly_agg_for_zone(zone_id)
        return create_hourly_charging_plan(
            grid_zone_id=zone_id,
            expected_daily_kwh=expected_daily_kwh,
            has_pv=has_pv,
            pv_kwp=pv_kwp,
            hourly_agg=hourly_agg,
            available_for_ev_kw=30.0,
        )

    def test_output_is_24_row_dataframe(self):
        plan = self._plan()
        assert isinstance(plan, pd.DataFrame), "Result should be a DataFrame"
        assert len(plan) == 24, f"Expected 24 rows, got {len(plan)}"

    def test_peak_hours_18_21_have_zero_charging(self):
        plan = self._plan(expected_daily_kwh=50.0)
        peak_rows = plan[plan["hour"].isin([18, 19, 20, 21])]
        for _, row in peak_rows.iterrows():
            assert row["recommended_charging_kw"] == 0.0, (
                f"Hour {row['hour']} is a peak hour but has charging: "
                f"{row['recommended_charging_kw']} kW"
            )

    def test_total_allocated_le_expected_daily_kwh(self):
        expected = 50.0
        plan = self._plan(expected_daily_kwh=expected)
        total = plan["recommended_charging_kw"].sum()
        assert total <= expected + 0.01, (
            f"Allocated {total:.2f} kWh exceeds expected {expected} kWh"
        )

    def test_zero_kwh_produces_all_zeros(self):
        plan = self._plan(expected_daily_kwh=0.0)
        assert plan["recommended_charging_kw"].sum() == 0.0

    def test_required_columns_present(self):
        plan = self._plan()
        for col in ["hour", "grid_score", "pv_generation_kw", "recommended_charging_kw", "reason"]:
            assert col in plan.columns, f"Missing column: {col}"

    def test_hours_are_0_to_23(self):
        plan = self._plan()
        assert list(plan["hour"]) == list(range(24))

    def test_pv_enabled_midday_bonus_reflected_in_scores(self):
        """With PV enabled, midday hours (10-14) should have higher grid_score
        than the same plan without PV."""
        plan_no_pv = self._plan(has_pv=False)
        plan_pv = self._plan(has_pv=True, pv_kwp=10.0)

        midday_hours = [10, 11, 12, 13, 14]
        score_no_pv = plan_no_pv[plan_no_pv["hour"].isin(midday_hours)]["grid_score"].mean()
        score_pv = plan_pv[plan_pv["hour"].isin(midday_hours)]["grid_score"].mean()

        assert score_pv > score_no_pv, (
            f"PV should raise midday grid scores. "
            f"No-PV avg={score_no_pv:.1f}, PV avg={score_pv:.1f}"
        )

    def test_pv_enabled_midday_preferred_for_charging(self):
        """With PV enabled and a small daily kWh target, midday hours should
        receive charging allocation (not just nights)."""
        plan = self._plan(has_pv=True, pv_kwp=10.0, expected_daily_kwh=20.0)
        midday = plan[plan["hour"].isin([10, 11, 12, 13, 14])]
        assert midday["recommended_charging_kw"].sum() > 0, (
            "With PV enabled, some midday charging should be scheduled"
        )

    def test_fallback_zone_works_when_no_zone_data(self):
        """Scheduler should use fallback flat profile for unknown zone IDs."""
        from scheduler import create_hourly_charging_plan
        empty_agg = pd.DataFrame(columns=["grid_zone_id", "hour", "avg_available_kw", "overload_rate"])
        plan = create_hourly_charging_plan(
            grid_zone_id="NONEXISTENT_ZONE",
            expected_daily_kwh=10.0,
            has_pv=False,
            pv_kwp=0.0,
            hourly_agg=empty_agg,
            available_for_ev_kw=20.0,
        )
        assert len(plan) == 24


# ===========================================================================
# 4. train_models
# ===========================================================================

def _make_capacity_result(available_for_ev_kw=20.0):
    """Minimal capacity_result dict for train_models tests."""
    return {
        "available_for_ev_kw": available_for_ev_kw,
        "max_wallboxes_without_balancing": 1,
        "max_simultaneous_with_balancing": 5,
        "total_building_capacity_kw": 41.5,
    }


class TestTrainModels:
    def test_model_loads_without_error(self):
        from train_models import load_or_train_classifier
        result = load_or_train_classifier()
        assert result is not None

    def test_model_tuple_has_two_classifiers(self):
        """load_or_train_classifier returns a (solution_clf, lb_clf) tuple."""
        from train_models import load_or_train_classifier
        result = load_or_train_classifier()
        assert len(result) == 2, f"Expected 2-tuple of classifiers, got {type(result)}"

    def test_classifiers_have_predict_method(self):
        from train_models import load_or_train_classifier
        sol_clf, lb_clf = load_or_train_classifier()
        assert hasattr(sol_clf, "predict"), "solution classifier should have predict method"
        assert hasattr(lb_clf, "predict"), "lb classifier should have predict method"

    def test_predict_solution_returns_dict_with_solution_type(self):
        from train_models import load_or_train_classifier, predict_solution
        model_tuple = load_or_train_classifier()
        zone_row = _make_zone_row()
        cap = _make_capacity_result()
        user_input = {"expected_evs": 5, "parking_spaces": 12, "flats": 24}
        result = predict_solution(model_tuple, zone_row, cap, user_input)
        assert isinstance(result, dict), f"Expected dict, got {type(result)}"
        assert "solution_type" in result, f"Missing 'solution_type' key, got {result}"

    def test_predict_solution_returns_dict_with_load_balancing_type(self):
        from train_models import load_or_train_classifier, predict_solution
        model_tuple = load_or_train_classifier()
        zone_row = _make_zone_row()
        cap = _make_capacity_result()
        user_input = {"expected_evs": 5, "parking_spaces": 12, "flats": 24}
        result = predict_solution(model_tuple, zone_row, cap, user_input)
        assert "load_balancing_type" in result, f"Missing 'load_balancing_type' key, got {result}"

    def test_predict_solution_returns_known_solution_label(self):
        from train_models import load_or_train_classifier, predict_solution
        from recommender import SOLUTION_LABELS
        model_tuple = load_or_train_classifier()
        zone_row = _make_zone_row()
        cap = _make_capacity_result()
        user_input = {"expected_evs": 5, "parking_spaces": 12, "flats": 24}
        result = predict_solution(model_tuple, zone_row, cap, user_input)
        assert result["solution_type"] in SOLUTION_LABELS, (
            f"Prediction '{result['solution_type']}' is not a known solution label: "
            f"{list(SOLUTION_LABELS.keys())}"
        )

    # -----------------------------------------------------------------------
    # CRITICAL: Expected to FAIL.
    # The model was trained without building parameters (available_for_ev_kw,
    # flats, etc. are in ALL_FEATURES but NOT in the training data), so passing
    # radically different capacity_result values returns the same prediction.
    # -----------------------------------------------------------------------
    def test_predictions_change_with_building_parameters_KNOWN_FAIL(self):
        """
        KNOWN FAILING TEST.

        predict_solution() now accepts capacity_result and user_input, and
        ALL_FEATURES includes building params (available_for_ev_kw, flats, etc.).
        However, the RandomForest was trained on zones_train.csv which contains
        ONLY zone-level columns — the building-param columns were all zero at
        training time, so the model learned to ignore them.

        As a result, passing a 0 kW available capacity vs 100 kW available
        capacity through predict_solution() should produce different predictions
        (none_monitor vs a real solution) but does NOT.
        """
        from train_models import load_or_train_classifier, predict_solution

        model_tuple = load_or_train_classifier()
        zone_row = _make_zone_row(reserve_margin=25.0, grid_sensitivity=0.4)
        user_input = {"expected_evs": 5, "parking_spaces": 12, "flats": 24}

        # No capacity at all — should suggest none_monitor
        cap_zero = _make_capacity_result(available_for_ev_kw=0.0)
        cap_zero["max_wallboxes_without_balancing"] = 0
        cap_zero["max_simultaneous_with_balancing"] = 0
        cap_zero["total_building_capacity_kw"] = 0.0

        # Ample capacity — should suggest a real wallbox solution
        cap_full = _make_capacity_result(available_for_ev_kw=100.0)
        cap_full["max_wallboxes_without_balancing"] = 9
        cap_full["max_simultaneous_with_balancing"] = 27
        cap_full["total_building_capacity_kw"] = 105.0

        pred_zero = predict_solution(model_tuple, zone_row, cap_zero, user_input)
        pred_full = predict_solution(model_tuple, zone_row, cap_full, user_input)

        assert pred_zero["solution_type"] != pred_full["solution_type"], (
            "KNOWN FAIL: ML predict_solution() returns identical solution_type "
            f"for 0 kW available ({pred_zero['solution_type']}) and 100 kW available "
            f"({pred_full['solution_type']}). "
            "The model was trained without building-parameter features, so it ignores them."
        )


# ===========================================================================
# 5. document_generator
# ===========================================================================

class TestDocumentGenerator:
    def _make_inputs(self):
        user_input = {
            "flats": 24,
            "parking_spaces": 12,
            "breaker_amps": 63,
            "expected_evs": 5,
            "has_pv": False,
            "pv_kwp": 0,
        }
        capacity_result = {
            "total_building_capacity_kw": 41.5,
            "estimated_evening_base_load_kw": 33.6,
            "available_for_ev_kw": 7.9,
            "max_wallboxes_without_balancing": 0,
            "max_simultaneous_with_balancing": 2,
            "recommended_installed_wallboxes": 2,
            "pv_peak_kw": 0.0,
            "warning_level": "orange",
        }
        recommendation = {
            "recommended_solution_type": "residential_ac_small",
            "solution_label": "Malá rezidenční AC stanice (2× 11 kW)",
            "number_of_wallboxes": 2,
            "recommended_total_kw": 22.0,
            "load_balancing_type": "static",
            "load_balancing_label": "Statický load balancing",
            "explanation": "Test explanation.",
            "risk_score": 20,
        }
        economics = {
            "base_cost": 50_000,
            "wallbox_cost": 40_000,
            "lb_cost": 0,
            "pv_cost": 0,
            "total_cost": 90_000,
            "monthly_savings": 2_000,
            "annual_savings": 24_000,
            "payback_years": 3.75,
        }
        return user_input, capacity_result, recommendation, economics

    def test_returns_bytesio(self):
        from document_generator import generate_svj_document
        user_input, capacity_result, recommendation, economics = self._make_inputs()
        result = generate_svj_document(user_input, capacity_result, recommendation, economics, "TS_TEST")
        assert isinstance(result, BytesIO), f"Expected BytesIO, got {type(result)}"

    def test_returns_non_empty_bytes(self):
        from document_generator import generate_svj_document
        user_input, capacity_result, recommendation, economics = self._make_inputs()
        result = generate_svj_document(user_input, capacity_result, recommendation, economics, "TS_TEST")
        data = result.read()
        assert len(data) > 0, "Document should not be empty"

    def test_can_be_opened_as_valid_docx(self):
        from document_generator import generate_svj_document
        from docx import Document
        user_input, capacity_result, recommendation, economics = self._make_inputs()
        buf = generate_svj_document(user_input, capacity_result, recommendation, economics, "TS_TEST")
        # Should not raise
        doc = Document(buf)
        assert doc is not None

    def test_docx_contains_expected_headings(self):
        from document_generator import generate_svj_document
        from docx import Document
        user_input, capacity_result, recommendation, economics = self._make_inputs()
        buf = generate_svj_document(user_input, capacity_result, recommendation, economics, "TS_TEST")
        doc = Document(buf)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Voltík" in full_text, "Document should contain 'Voltík'"
        assert "SVJ" in full_text, "Document should contain 'SVJ'"

    def test_docx_with_pv_includes_pv_info(self):
        """When PV is enabled the document should mention it in the building params table."""
        from document_generator import generate_svj_document
        from docx import Document
        user_input, capacity_result, recommendation, economics = self._make_inputs()
        user_input["has_pv"] = True
        user_input["pv_kwp"] = 15.0
        buf = generate_svj_document(user_input, capacity_result, recommendation, economics, "TS_TEST")
        doc = Document(buf)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        # Tables text comes via paragraphs inside cells; also check table cells
        table_text = ""
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    table_text += cell.text + "\n"
        combined = full_text + table_text
        assert "15" in combined, "Document should reflect the PV kWp value"
