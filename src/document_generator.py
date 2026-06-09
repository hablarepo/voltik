from datetime import date
from io import BytesIO
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def _row(table, label: str, value: str):
    row = table.add_row()
    row.cells[0].text = label
    row.cells[1].text = value


def generate_svj_document(
    user_input: dict,
    capacity_result: dict,
    recommendation: dict,
    economics: dict,
    zone_id: str,
) -> BytesIO:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Title
    title = doc.add_heading("Voltík – Podklad pro hlasování SVJ", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"Datum zpracování: {date.today().strftime('%d. %m. %Y')}")
    doc.add_paragraph(
        "⚠️ Tento dokument je demonstrační odhad vytvořený nástrojem Voltík. "
        "Není náhradou za odborný elektroprojekt ani závazným inženýrským posudkem."
    ).runs[0].italic = True
    doc.add_paragraph()

    # Building parameters
    _heading(doc, "1. Parametry bytového domu", 1)
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.rows[0].cells[0].text = "Parametr"
    t.rows[0].cells[1].text = "Hodnota"
    _row(t, "Počet bytových jednotek", str(user_input["flats"]))
    _row(t, "Počet parkovacích míst", str(user_input["parking_spaces"]))
    _row(t, "Hlavní jistič", f"3×{user_input['breaker_amps']} A")
    _row(t, "Očekávaný počet EV", str(user_input["expected_evs"]))
    _row(t, "Fotovoltaika", f"Ano – {user_input.get('pv_kwp', 0)} kWp" if user_input.get("has_pv") else "Ne")
    _row(t, "Distribuční zóna Praha", zone_id)
    doc.add_paragraph()

    # Capacity
    _heading(doc, "2. Kapacitní analýza", 1)
    t2 = doc.add_table(rows=1, cols=2)
    t2.style = "Table Grid"
    t2.rows[0].cells[0].text = "Ukazatel"
    t2.rows[0].cells[1].text = "Hodnota"
    _row(t2, "Celková kapacita přípojky", f"{capacity_result['total_building_capacity_kw']} kW")
    _row(t2, "Odhadovaná večerní spotřeba domu", f"{capacity_result['estimated_evening_base_load_kw']} kW")
    _row(t2, "Dostupný výkon pro EV", f"{capacity_result['available_for_ev_kw']} kW")
    _row(t2, "Max. wallboxy bez load balancingu", str(capacity_result["max_wallboxes_without_balancing"]))
    _row(t2, "Max. wallboxy s dynamickým LB", str(capacity_result["max_simultaneous_with_balancing"]))
    doc.add_paragraph()

    # Recommendation
    _heading(doc, "3. Doporučená konfigurace", 1)
    t3 = doc.add_table(rows=1, cols=2)
    t3.style = "Table Grid"
    t3.rows[0].cells[0].text = "Parametr"
    t3.rows[0].cells[1].text = "Hodnota"
    _row(t3, "Typ řešení", recommendation["solution_label"])
    _row(t3, "Počet wallboxů", str(recommendation["number_of_wallboxes"]))
    _row(t3, "Celkový instalovaný výkon", f"{recommendation['recommended_total_kw']:.0f} kW")
    _row(t3, "Load balancing", recommendation["load_balancing_label"])
    doc.add_paragraph()

    # Cost estimate
    _heading(doc, "4. Odhadované náklady", 1)
    t4 = doc.add_table(rows=1, cols=2)
    t4.style = "Table Grid"
    t4.rows[0].cells[0].text = "Položka"
    t4.rows[0].cells[1].text = "Odhadovaná cena"
    _row(t4, "Základní projekt a elektroinstalace", f"{economics['base_cost']:,.0f} Kč")
    _row(t4, f"Wallboxy ({recommendation['number_of_wallboxes']}×)", f"{economics['wallbox_cost']:,.0f} Kč")
    if economics.get("lb_cost", 0) > 0:
        _row(t4, "Dynamický load balancing systém", f"{economics['lb_cost']:,.0f} Kč")
    if economics.get("pv_cost", 0) > 0:
        _row(t4, "Integrace FVE", f"{economics['pv_cost']:,.0f} Kč")
    _row(t4, "Celkem (odhad)", f"{economics['total_cost']:,.0f} Kč")
    doc.add_paragraph()

    # Savings
    _heading(doc, "5. Ekonomická úspora", 1)
    t5 = doc.add_table(rows=1, cols=2)
    t5.style = "Table Grid"
    t5.rows[0].cells[0].text = "Ukazatel"
    t5.rows[0].cells[1].text = "Hodnota"
    _row(t5, "Odhadovaná měsíční úspora vs. veřejné nabíjení", f"{economics['monthly_savings']:,.0f} Kč/měsíc")
    _row(t5, "Odhadovaná roční úspora", f"{economics['annual_savings']:,.0f} Kč/rok")
    _row(t5, "Orientační návratnost", f"{economics['payback_years']:.1f} let")
    doc.add_paragraph()

    # Cost sharing
    _heading(doc, "6. Návrh rozdělení nákladů", 1)
    doc.add_paragraph(
        "Doporučujeme zvážit následující modely financování:\n"
        "• Individuální příspěvek: každý vlastník s wallboxem hradí poměrnou část nákladů.\n"
        "• Fond oprav SVJ: investice z fondu oprav, wallboxy jako společné zařízení.\n"
        "• Dotační programy: aktuálně dostupné dotace (NPO, OP TAK) mohou pokrýt část nákladů.\n"
        "Konkrétní způsob schvaluje shromáždění vlastníků."
    )
    doc.add_paragraph()

    # Resolution text
    _heading(doc, "7. Návrh usnesení shromáždění vlastníků", 1)
    res = doc.add_paragraph()
    res.add_run(
        "Shromáždění vlastníků schvaluje přípravu a realizaci dobíjecí infrastruktury "
        f"pro elektrická vozidla v bytovém domě. Instalace bude zahrnovat {recommendation['number_of_wallboxes']} "
        f"wallbox(ů) s celkovým výkonem {recommendation['recommended_total_kw']:.0f} kW "
        f"a {recommendation['load_balancing_label'].lower()}. "
        "SVJ pověřuje výbor zajištěním elektroprojektové dokumentace, výběrem dodavatele "
        "a podáním případné žádosti o dotaci. Celkové náklady nesmí přesáhnout "
        f"{economics['total_cost'] * 1.15:,.0f} Kč (včetně 15% rezervy) bez dalšího souhlasu shromáždění."
    ).bold = False
    doc.add_paragraph()

    # Risk note
    _heading(doc, "8. Rizika a poznámky ke spravedlivosti", 1)
    doc.add_paragraph(
        "• Reálná kapacita přípojky musí být ověřena revizním technikem a distributorem (PREdistribuce).\n"
        "• Dynamický load balancing zajišťuje spravedlivé sdílení výkonu mezi wallboxy.\n"
        "• Vlastníci bez vozidla EV mohou v budoucnu požádat o přidání wallboxu na vlastní náklady.\n"
        "• Přetížení sítě v distribuční zóně může vyžadovat navýšení jisticího příkonu – koordinovat s PRE."
    )
    doc.add_paragraph()

    # Disclaimer
    _heading(doc, "Prohlášení", 1)
    disc = doc.add_paragraph(
        "Tento dokument byl vygenerován nástrojem Voltík (hackathonový demonstrátor) na základě "
        "uživatelem zadaných parametrů a dat o pražské distribuční síti. Nejde o závazný technický "
        "posudek ani o elektroprojektovou dokumentaci. Před realizací je nutné zpracovat odborný projekt "
        "a získat souhlas distributora elektrické energie."
    )
    disc.runs[0].italic = True if disc.runs else None

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
